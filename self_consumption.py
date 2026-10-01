"""Profili orari e bilancio FV/utenza. Nessuna batteria nel calcolo."""
from __future__ import annotations

import csv
import datetime as dt
import json
import math
from pathlib import Path
from zoneinfo import ZoneInfo


def orientation_key(orientation):
    return ','.join(f'{float(v):g}' for v in orientation)


def _load_weather_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def is_pvgis_tmy_json(data):
    """Return True for the PVGIS TMY JSON structure."""
    return (
        isinstance(data, dict)
        and isinstance(data.get('inputs'), dict)
        and isinstance(data.get('outputs'), dict)
        and isinstance(data['outputs'].get('tmy_hourly'), list)
    )


def _pvgis_stamp(value):
    """Parse PVGIS timestamps such as 20220101:1300 as UTC datetimes."""
    try:
        return dt.datetime.strptime(str(value), '%Y%m%d:%H%M').replace(tzinfo=dt.timezone.utc)
    except (TypeError, ValueError) as exc:
        raise ValueError(f'Invalid PVGIS TMY timestamp: {value!r}') from exc


def _target_utc_hours(profile, timezone):
    zone = ZoneInfo(timezone)
    result = []
    for stamp in timestamps(profile):
        aware = stamp.replace(tzinfo=zone) if stamp.tzinfo is None else stamp.astimezone(zone)
        result.append(aware.astimezone(dt.timezone.utc))
    return result


def _pvgis_row_values(row):
    keys = ('T2m', 'G(h)', 'Gb(n)', 'Gd(h)')
    try:
        values = {key: float(row[key]) for key in keys}
    except (KeyError, TypeError, ValueError) as exc:
        raise ValueError('PVGIS TMY must contain T2m, G(h), Gb(n) and Gd(h) for every hour') from exc
    if not all(math.isfinite(v) for v in values.values()):
        raise ValueError('PVGIS TMY contains a non-finite weather value')
    if not -60 <= values['T2m'] <= 65:
        raise ValueError('PVGIS TMY temperature is outside the accepted range')
    if any(values[k] < 0 or values[k] > 1600 for k in ('G(h)', 'Gb(n)', 'Gd(h)')):
        raise ValueError('PVGIS TMY irradiance is outside the accepted range')
    return values


def _pvgis_lookup(data):
    rows = data.get('outputs', {}).get('tmy_hourly', [])
    if not rows:
        raise ValueError('PVGIS JSON does not contain outputs.tmy_hourly')
    lookup = {}
    for row in rows:
        stamp = _pvgis_stamp(row.get('time(UTC)'))
        key = (stamp.month, stamp.day, stamp.hour)
        if key in lookup:
            raise ValueError(f'Duplicate PVGIS TMY hour: {key}')
        lookup[key] = _pvgis_row_values(row)
    if len(lookup) < 8760:
        raise ValueError(f'PVGIS TMY is incomplete: {len(lookup)} hourly records found')
    return lookup


def _lookup_tmy_hour(lookup, stamp_utc):
    key = (stamp_utc.month, stamp_utc.day, stamp_utc.hour)
    if key in lookup:
        return lookup[key]
    if stamp_utc.month == 2 and stamp_utc.day == 29:
        a = lookup.get((2, 28, stamp_utc.hour))
        b = lookup.get((3, 1, stamp_utc.hour))
        if a and b:
            return {k: (a[k] + b[k]) / 2 for k in a}
    raise ValueError(f'PVGIS TMY does not contain {stamp_utc:%m-%d %H}:00 UTC')


def _poa_isotropic(ghi, dni, dhi, elevation_deg, solar_azimuth_deg,
                   tilt_deg, panel_azimuth_deg, albedo=0.20):
    """Plane-of-array irradiance from DNI/DHI/GHI using an isotropic-sky model."""
    if elevation_deg <= 0:
        return 0.0
    elev = math.radians(elevation_deg)
    beta = math.radians(float(tilt_deg))
    delta_az = math.radians(float(solar_azimuth_deg) - float(panel_azimuth_deg))
    cos_incidence = (
        math.sin(elev) * math.cos(beta)
        + math.cos(elev) * math.sin(beta) * math.cos(delta_az)
    )
    beam = max(0.0, float(dni) * max(0.0, cos_incidence))
    diffuse = max(0.0, float(dhi)) * (1.0 + math.cos(beta)) / 2.0
    ground = max(0.0, float(ghi)) * float(albedo) * (1.0 - math.cos(beta)) / 2.0
    return max(0.0, beam + diffuse + ground)


def read_pvgis_tmy_json(path, profile, orientations, solar_position, timezone='Europe/Rome',
                        albedo=0.20):
    """Convert a PVGIS TMY JSON into the application's per-orientation weather profiles."""
    data = _load_weather_json(path)
    if not is_pvgis_tmy_json(data):
        raise ValueError('Not a PVGIS TMY JSON file')
    location = data.get('inputs', {}).get('location', {})
    try:
        latitude = float(location['latitude'])
        longitude = float(location['longitude'])
    except (KeyError, TypeError, ValueError) as exc:
        raise ValueError('PVGIS TMY location metadata is missing') from exc
    if not all(math.isfinite(v) for v in (latitude, longitude)):
        raise ValueError('PVGIS TMY location metadata is invalid')

    lookup = _pvgis_lookup(data)
    target_utc = _target_utc_hours(profile, timezone)
    rows = [_lookup_tmy_hour(lookup, stamp) for stamp in target_utc]
    ambient = [row['T2m'] for row in rows]
    zone = ZoneInfo(timezone)

    unique_orientations = []
    for orientation in orientations:
        pair = (float(orientation[0]), float(orientation[1]) % 360)
        if pair not in unique_orientations:
            unique_orientations.append(pair)
    if not unique_orientations:
        raise ValueError('No panel orientation is available for the PVGIS TMY conversion')

    result = {}
    for tilt, azimuth in unique_orientations:
        poa = []
        for stamp_utc, row in zip(target_utc, rows):
            midpoint_utc = stamp_utc + dt.timedelta(minutes=30)
            midpoint = midpoint_utc.astimezone(zone)
            offset = midpoint.utcoffset().total_seconds() / 3600
            elevation, solar_azimuth = solar_position(
                latitude, longitude, midpoint.day, midpoint.month,
                midpoint.hour + midpoint.minute / 60 + midpoint.second / 3600,
                offset, year=midpoint.year
            )
            poa.append(_poa_isotropic(
                row['G(h)'], row['Gb(n)'], row['Gd(h)'],
                elevation, solar_azimuth, tilt, azimuth, albedo
            ))
        result[orientation_key((tilt, azimuth))] = {
            'hourly_poa_w_m2': poa,
            'hourly_ambient_c': list(ambient),
            'source_filename': Path(path).name,
            'source': 'PVGIS typical meteorological year',
            'period_start': profile['start_date'],
            'period_end': profile['end_date'],
            'latitude': latitude,
            'longitude': longitude,
            'query': {
                'latitude': latitude,
                'longitude': longitude,
                'tilt_deg': tilt,
                'azimuth_deg_from_south': (azimuth - 180 + 180) % 360 - 180,
            },
            'tmy': True,
            'tmy_year_min': data.get('inputs', {}).get('meteo_data', {}).get('year_min'),
            'tmy_year_max': data.get('inputs', {}).get('meteo_data', {}).get('year_max'),
            'tmy_months_selected': data.get('outputs', {}).get('months_selected', []),
        }
    return {'orientations': result, 'source': 'PVGIS TMY',
            'latitude': latitude, 'longitude': longitude}


def read_weather_json(path, profile, timezone='Europe/Rome', orientations=None, solar_position=None):
    """Read either the existing Open-Meteo JSON or a PVGIS TMY JSON."""
    data = _load_weather_json(path)
    if is_pvgis_tmy_json(data):
        if orientations is None or solar_position is None:
            raise ValueError('PVGIS TMY import requires project panel orientations')
        return read_pvgis_tmy_json(
            path, profile, orientations, solar_position, timezone=timezone
        )
    return read_open_meteo_json(path, profile, timezone=timezone)


def read_open_meteo_json(path, profile, timezone="Europe/Rome"):
    """Allinea meteo storico Open-Meteo alle 24 colonne locali per giorno."""
    data=_load_weather_json(path)
    hourly=data.get('hourly',{})
    times=hourly.get('time',[])
    irradiance=hourly.get('global_tilted_irradiance',[])
    ambient=hourly.get('temperature_2m',[])
    if not times or len(times)!=len(irradiance) or len(times)!=len(ambient):
        raise ValueError('JSON must contain hourly.time, global_tilted_irradiance and temperature_2m')
    if data.get('timezone') != timezone:
        raise ValueError(f'The weather file must use timezone={timezone}')
    groups={}
    repeated={}
    for stamp,poa,temp in zip(times,irradiance,ambient):
        # Open-Meteo documenta GTI come media dell'ora precedente: etichetta
        # l'intervallo all'inizio, coerente con la colonna 0..23 dei consumi.
        source=dt.datetime.fromisoformat(stamp)
        if profile.get('timestamps'):
            if source.tzinfo is None:
                fold=repeated.get(stamp,0);repeated[stamp]=fold+1
                source=source.replace(tzinfo=ZoneInfo(timezone),fold=min(fold,1))
            key=(source.astimezone(dt.timezone.utc)-dt.timedelta(hours=1)).isoformat()
        else:key=(source-dt.timedelta(hours=1)).strftime('%Y-%m-%dT%H:00')
        if poa is None or temp is None or not all(isinstance(x,(int,float)) and math.isfinite(x) for x in (poa,temp)):
            raise ValueError(f'Missing or invalid weather data: {stamp}')
        if not 0 <= poa <= 1600 or not -60 <= temp <= 65:
            raise ValueError(f'Weather value outside the accepted range: {stamp}')
        groups.setdefault(key,[]).append((float(poa),float(temp)))
    hourly_poa=[];hourly_temp=[];ambiguous=[];absent=[]
    stamps=timestamps(profile)
    for i,stamp in enumerate(stamps):
        key=stamp.astimezone(dt.timezone.utc).isoformat() if profile.get('timestamps') else stamp.strftime('%Y-%m-%dT%H:00')
        pairs=groups.get(key)
        if not pairs:
            hourly_poa.append(None);hourly_temp.append(None);absent.append(i)
        else:
            if len(pairs)>1:ambiguous.append(i)
            hourly_poa.append(sum(p[0] for p in pairs)/len(pairs))
            hourly_temp.append(sum(p[1] for p in pairs)/len(pairs))
    if len(absent)>1 or len(ambiguous)>1:
        raise ValueError(f'Weather timestamps do not match: {len(absent)} missing, {len(ambiguous)} duplicated')
    if profile.get('timestamps') and (absent or ambiguous):raise ValueError('Exact timestamp profiles require one matching weather value per hour.')
    for i in absent:
        if i%24==0 or i%24==23:
            raise ValueError('Cannot interpolate a missing hour at the start or end of a day')
        for seq in (hourly_poa,hourly_temp):
            if seq[i-1] is None or seq[i+1] is None:
                raise ValueError('Cannot interpolate weather: neighbouring hours are missing')
            seq[i]=(seq[i-1]+seq[i+1])/2
    return {'hourly_poa_w_m2':hourly_poa,'hourly_ambient_c':hourly_temp,
            'source_filename':Path(path).name,'source':'Open-Meteo historical reanalysis',
            'missing_local_indices':absent,'repeated_local_indices':ambiguous,
            'period_start':profile['start_date'],'period_end':profile['end_date'],
            'latitude':data.get('latitude'),'longitude':data.get('longitude')}


def read_daily_excel(path, imputation='previous_week'):
    """Legge il formato Volfrigo: data, giorno, colonne 0..23 in kWh."""
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        header = next((row for row in rows if row and str(row[0]).strip().lower() in ('date', 'data')), None)
        if header is None or [str(x) for x in header[2:26]] != [str(h) for h in range(24)]:
            raise ValueError('Expected columns: date, day, 0, ..., 23 (legacy data/giorno headers also accepted)')
        dates, values, missing = [], [], []
        for row in rows:
            if not row or not isinstance(row[0], (dt.date, dt.datetime)):
                continue
            date = row[0].date() if isinstance(row[0], dt.datetime) else row[0]
            if dates and date != dates[-1] + dt.timedelta(days=1):
                raise ValueError(f'Dates are not consecutive: {dates[-1]} / {date}')
            dates.append(date)
            for h in range(24):
                value = row[h+2] if len(row) > h+2 else None
                if value is None or value == '':
                    missing.append(len(values)); values.append(None)
                elif isinstance(value, (int, float)) and math.isfinite(value) and value >= 0:
                    values.append(float(value))
                else:
                    raise ValueError(f'Invalid kWh value: {date}, hour {h}: {value!r}')
        if not dates:
            raise ValueError('No hourly consumption values found')
        if len(values) != len(dates)*24:
            raise ValueError('The number of hourly values is inconsistent')
        # Prior 7 same clock hours: preserves a 24-slot/day representation.
        if missing and imputation=='reject':raise ValueError('Consumption contains missing values; fill them or select an imputation method.')
        if imputation not in ('previous_week','reject','zero'):raise ValueError('Unknown consumption imputation method.')
        original_values=list(values)
        for i in missing:
            if imputation=='zero':values[i]=0.;continue
            history = [original_values[i-24*d] for d in range(1, 8)
                       if i-24*d >= 0 and original_values[i-24*d] is not None]
            if not history:
                raise ValueError(f'Cannot estimate missing consumption at index {i}')
            values[i] = sum(history)/len(history)
        return {'start_date': dates[0].isoformat(), 'end_date': dates[-1].isoformat(),
                'hourly_kwh': values, 'imputed_indices': missing,
                'source_filename': Path(path).name, 'imputation_method':imputation, 'hour_convention': 'legacy 24 slots/day; DST approximate'}
    finally:
        book.close()


def timestamps(profile):
    if profile.get('timestamps'):
        result=[dt.datetime.fromisoformat(stamp) for stamp in profile['timestamps']]
        if len(result)!=len(profile['hourly_kwh']):raise ValueError('Timestamp and load counts differ.')
        utc=[stamp.astimezone(dt.timezone.utc) for stamp in result]
        if any(stamp.tzinfo is None for stamp in result) or any((b-a).total_seconds()!=3600 for a,b in zip(utc,utc[1:])):
            raise ValueError('Use consecutive hourly timestamps with explicit UTC offsets.')
        return result
    start = dt.date.fromisoformat(profile['start_date'])
    end = dt.date.fromisoformat(profile['end_date'])
    n = (end-start).days+1
    if n*24 != len(profile['hourly_kwh']):
        raise ValueError('Consumption profile does not match the specified dates')
    return [dt.datetime.combine(start + dt.timedelta(days=i//24),dt.time(i%24))
            for i in range(n*24)]


def production_from_program(app, profile, ac_factor=0.90, progress=None, weather=None):
    """Usa gli stessi metodi di energia, sole e ombra del programma.

    Il simulatore originale è a cielo sereno; ac_factor modellizza perdite
    aggiuntive DC/AC in uno scenario parametrico, non dati meteo misurati.
    """
    if not 0 < ac_factor <= 1:
        raise ValueError('The additional AC factor must be greater than zero and at most one')
    if not app.panels or getattr(app, 'panel_pmax_w', 0) <= 0:
        raise ValueError('Load a panel layout with a valid module Pmax rating')
    if getattr(app,'pylon_img_pos',None) is not None and getattr(app,'px_per_mm',0) <= 0:
        raise ValueError('Calibrate the image scale before calculating obstacle shading')
    zone = ZoneInfo(getattr(app,'model_settings',{}).get('timezone','Europe/Rome'))
    from engineering_inputs import panel_connections, ac_limit_kw, module_for_string
    connections=panel_connections(app)
    project=app._project_snapshot()
    ac_limits={block:ac_limit_kw(project,block) for block in set(connections.values())}
    panel_ratings={}
    for sid,members in app.strings.items():
        spec=module_for_string(project,sid)
        if not spec and len(project.get('material_categories',{}).get('modules',{}).get('rows',[]))>1:
            raise ValueError(f'{sid}: select its module equipment row in String sizing.')
        rating=spec.get('Pmax (W)',app.panel_pmax_w)
        try:rating=float(str(rating).replace(',','.'))
        except (TypeError,ValueError):raise ValueError(f'{sid}: enter module Pmax.')
        if not math.isfinite(rating) or rating<=0:raise ValueError(f'{sid}: invalid module Pmax.')
        for coord in members:panel_ratings[coord]=rating
    orientations={}
    coord_group={}
    for coord in app.panels:
        item=app.panel_orientations.get(coord,{})
        key=(item.get('tilt_deg',app.panel_tilt_deg),item.get('azimuth_deg',app.panel_azimuth_deg))
        orientations.setdefault(key,[]).append(coord);coord_group[coord]=key
    weather_by_orientation={}
    if weather:
        if 'orientations' in weather:weather_by_orientation=weather['orientations']
        elif len(orientations)==1:weather_by_orientation={orientation_key(next(iter(orientations))):weather}
        else:raise ValueError('Import one historical GTI profile per orientation using Import weather JSON.')
        for orientation in orientations:
            item=weather_by_orientation.get(orientation_key(orientation))
            if not item:raise ValueError(f'Missing weather profile for orientation {orientation}.')
            query=item.get('query',{})
            if not query:raise ValueError('Reimport weather with its source location and orientation before use.')
            comparisons={'latitude':app.solar_latitude,'longitude':app.solar_longitude,'tilt_deg':orientation[0],
                         'azimuth_deg_from_south':(orientation[1]-180+180)%360-180}
            location_tolerance = 0.05 if item.get('tmy') else 1e-6
            for key,value in comparisons.items():
                if key not in query:
                    raise ValueError('Weather location or orientation metadata is missing. Reimport the weather file.')
                tolerance = location_tolerance if key in ('latitude','longitude') else 1e-6
                if abs(float(query[key])-value) > tolerance:
                    raise ValueError('Weather location or orientation changed. Import/download a matching profile.')
    pv = []
    hours = timestamps(profile)
    for item in weather_by_orientation.values():
        if (len(item['hourly_poa_w_m2'])!=len(hours) or len(item['hourly_ambient_c'])!=len(hours)
            or item['period_start']!=profile['start_date'] or item['period_end']!=profile['end_date']):
            raise ValueError('Weather period or number of hours does not match consumption.')
    for i, stamp in enumerate(hours):
        if stamp.tzinfo is not None:stamp=stamp.astimezone(zone)
        # Il file assume 24 ore anche nelle due date del cambio di ora.
        # Per ciascuna colonna si usa l'offset locale del suo centro orario.
        midpoint=((stamp.astimezone(dt.timezone.utc)+dt.timedelta(minutes=30)).astimezone(zone)
                  if stamp.tzinfo is not None else (stamp+dt.timedelta(minutes=30)).replace(tzinfo=zone))
        utc_offset = midpoint.utcoffset().total_seconds()/3600
        elev, az = app._compute_solar_position(
            app.solar_latitude,app.solar_longitude,midpoint.day,midpoint.month,
            midpoint.hour+midpoint.minute/60+midpoint.second/3600,utc_offset,year=midpoint.year)
        kwh = 0.0
        if elev > 0.1:
            irradiances={key:(weather_by_orientation[orientation_key(key)]['hourly_poa_w_m2'][i] if weather else
                              app._get_clear_sky_poa_irradiance(elev,az,members[0]))
                         for key,members in orientations.items()}
            ambient = next(iter(weather_by_orientation.values()))['hourly_ambient_c'][i] if weather else None
            def panel_w(irradiance,shaded_fraction=0):
                if ambient is None:
                    return app._estimate_panel_power_w(irradiance,shaded_fraction)
                from shading_models import attenuation
                effective=max(0,irradiance*(1-attenuation(shaded_fraction,app.model_settings)))
                cell_c=ambient+(app.panel_noct_c-20)*effective/800
                return max(0,app.panel_pmax_w*effective/1000*
                           max(0,1+app.panel_temp_coeff_pct/100*(cell_c-25)))
            ideal_power = {key:panel_w(g) for key,g in irradiances.items()}
            powers={coord:ideal_power[coord_group[coord]]*panel_ratings[coord]/app.panel_pmax_w for coord in app.panels}
            if app.pylon_img_pos is not None or app.model_settings.get('shadow_obstacles'):
                percentages,_,_=app._calculate_shadow_percentages(elev,az)
                for coord,pct in percentages.items():
                    key=coord_group[coord]
                    powers[coord]=panel_w(irradiances[key],pct/100)*panel_ratings[coord]/app.panel_pmax_w
            per_block={block:0. for block in ac_limits}
            for coord,power in powers.items():per_block[connections[coord]]+=power*ac_factor/1000
            kwh=sum(min(ac_limits[block],power) for block,power in per_block.items())
        pv.append(kwh)
        if progress and (i%72 == 0 or i == len(hours)-1):
            progress(i+1,len(hours))
    return pv


def balance(profile, pv_kwh, export_limit_kw=None):
    if export_limit_kw is not None and (not math.isfinite(export_limit_kw) or export_limit_kw < 0):
        raise ValueError("Invalid grid export limit")
    hours = timestamps(profile)
    if len(pv_kwh) != len(hours):
        raise ValueError('PV production and consumption must cover the same hours')
    load = profile['hourly_kwh']
    monthly = {}
    rows = []
    for i, (stamp, demand, produced) in enumerate(zip(hours,load,pv_kwh)):
        if not all(isinstance(v,(int,float)) and math.isfinite(v) and v>=0 for v in (demand,produced)):
            raise ValueError(f'Hour {i}: invalid production or consumption')
        self_used = min(demand,produced)
        imported = demand-self_used
        surplus = produced-self_used
        exported = surplus if export_limit_kw is None else min(surplus, export_limit_kw)
        curtailed = surplus-exported
        month = stamp.strftime('%Y-%m')
        bucket = monthly.setdefault(month,{'load_kwh':0,'pv_kwh':0,'self_kwh':0,'grid_kwh':0,'export_kwh':0,'curtailed_kwh':0})
        for key,val in (('load_kwh',demand),('pv_kwh',produced),('self_kwh',self_used),
                        ('grid_kwh',imported),('export_kwh',exported),('curtailed_kwh',curtailed)):
            bucket[key]+=val
        rows.append((stamp.isoformat(' '),demand,produced,self_used,imported,exported,i in profile.get('imputed_indices',[])))
    total={k:sum(m[k] for m in monthly.values()) for k in next(iter(monthly.values()))}
    total['autoconsumption_pct']=100*total['self_kwh']/total['pv_kwh'] if total['pv_kwh'] else 0
    total['self_sufficiency_pct']=100*total['self_kwh']/total['load_kwh'] if total['load_kwh'] else 0
    return {'annual':total,'monthly':monthly,'rows':rows,'export_limit_kw':export_limit_kw}


def write_hourly_csv(path, result):
    with open(path,'w',newline='',encoding='utf-8-sig') as out:
        writer=csv.writer(out)
        writer.writerow(['local_datetime','load_kwh','pv_ac_kwh','self_consumption_kwh',
                         'grid_import_kwh','grid_export_kwh','estimated_load'])
        writer.writerows(result['rows'])


def read_hourly_csv(path,timezone=None):
    with open(path,encoding='utf-8-sig',newline='') as stream:
        rows=list(csv.DictReader(stream))
    if not rows:raise ValueError('Empty hourly file.')
    stamps=[row['timestamp'] for row in rows]
    if timezone:
        zone=ZoneInfo(timezone)
        parsed=[dt.datetime.fromisoformat(s) for s in stamps]
        if any(s.tzinfo is None for s in parsed):raise ValueError('CSV timestamps require explicit UTC offsets.')
        stamps=[s.astimezone(zone).isoformat() for s in parsed]
    values=[float(row['consumption_kwh']) for row in rows]
    if any(not math.isfinite(x) or x<0 for x in values):raise ValueError('Invalid hourly kWh value.')
    profile={'timestamps':stamps,'hourly_kwh':values,'imputed_indices':[],
             'start_date':dt.datetime.fromisoformat(stamps[0]).date().isoformat(),
             'end_date':dt.datetime.fromisoformat(stamps[-1]).date().isoformat(),
             'hour_convention':'explicit offset timestamps','source_filename':Path(path).name}
    timestamps(profile)
    return profile
